"""Read supported static Excel networks without changing source files.

Supported schemas, units and examples: README_network_loader.md.
Hourly profiles belong to data_loader.py, not this module.
"""
from math import isclose, isfinite
from pathlib import Path
import re
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .network_model import Bus, Line, Load, NetworkModel, PVUnit


def _literal(value, label):
    if isinstance(value, str) and value.startswith("="):
        raise ValueError(f"{label}: use a literal value, not an Excel formula")
    return value


def _number(value, label):
    _literal(value, label)
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not isfinite(value):
        raise ValueError(f"{label}: expected a finite numeric value")
    return float(value)


def _integer(value, label):
    number = _number(value, label)
    if not number.is_integer():
        raise ValueError(f"{label}: expected an integer")
    return int(value)


def _text(value, label):
    _literal(value, label)
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{label}: expected nonempty text")
    return value.strip()


def _flag(value, label):
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)) and value in (0, 1):
        return bool(value)
    raise ValueError(f"{label}: use TRUE/FALSE or numeric 1/0")


def _table(sheet, required, optional=(), reference_footer=False):
    rows = iter(sheet.iter_rows(values_only=True))
    header = next(rows, ())
    names = [value.strip() if isinstance(value, str) else value for value in header]
    populated = [value for value in names if value is not None]
    if any(not isinstance(value, str) or not value for value in populated):
        raise ValueError(f"{sheet.title}: header names must be nonempty text")
    if len(set(populated)) != len(populated):
        raise ValueError(f"{sheet.title}: duplicate header names")
    missing = set(required) - set(populated)
    unknown = set(populated) - set(required) - set(optional)
    if missing or unknown:
        raise ValueError(f"{sheet.title}: missing columns {sorted(missing)}; unsupported columns {sorted(unknown)}")
    records = []
    footer_seen = False
    for row_number, values in enumerate(rows, 2):
        if all(value is None for value in values):
            continue
        first = values[0] if values else None
        if reference_footer and isinstance(first, str):
            if first.strip() == "TOTAL":
                if footer_seen:
                    raise ValueError(f"{sheet.title}: repeated TOTAL footer")
                footer_seen = True
                continue
            if first.startswith("Source:") and all(value is None for value in values[1:]):
                footer_seen = True
                continue
        if footer_seen:
            raise ValueError(f"{sheet.title} row {row_number}: data after the footer")
        record = {}
        for index, name in enumerate(names):
            value = values[index] if index < len(values) else None
            if name is None and value is not None:
                raise ValueError(f"{sheet.title} row {row_number}: data in an unnamed column")
            if name is not None:
                record[name] = value
        records.append((row_number, record))
    return records


def _check_sheets(book, required, allowed):
    missing = set(required) - set(book.sheetnames)
    unknown = set(book.sheetnames) - set(allowed)
    if missing or unknown:
        raise ValueError(f"Workbook: missing sheets {sorted(missing)}; unsupported sheets {sorted(unknown)}")


def _project_network(book, frequency_hz):
    _check_sheets(book, ("Network", "Buses", "Lines"), ("Network", "Buses", "Lines", "Loads", "PV"))
    settings = {}
    for row, record in _table(book["Network"], ("parameter", "value")):
        name = _text(record["parameter"], f"Network row {row} parameter")
        if name not in ("base_power_kva", "frequency_hz", "name", "notes"):
            raise ValueError(f"Network row {row}: unsupported parameter {name}")
        if name in settings:
            raise ValueError(f"Network: duplicate parameter {name}")
        settings[name] = record["value"]
    base = _number(settings.get("base_power_kva"), "Network base_power_kva")
    if frequency_hz is None:
        frequency_hz = _number(settings.get("frequency_hz"), "Network frequency_hz")
    elif "frequency_hz" in settings:
        if _number(settings["frequency_hz"], "Network frequency_hz") <= 0:
            raise ValueError("Network frequency_hz must be positive")
    buses, lines, loads, pv_units = [], [], [], []
    for row, record in _table(book["Buses"], ("bus_id", "base_kv", "is_slack"),
                              ("name", "v_set_pu", "notes")):
        label = f"Buses row {row}"
        bus_id = _integer(record["bus_id"], f"{label} bus_id")
        name = _text(record["name"], f"{label} name") if record.get("name") is not None else f"Bus {bus_id}"
        voltage = _number(record["v_set_pu"], f"{label} v_set_pu") if record.get("v_set_pu") is not None else 1.0
        buses.append(Bus(bus_id, name, _number(record["base_kv"], f"{label} base_kv"),
                         _flag(record["is_slack"], f"{label} is_slack"), voltage))
    bus_ids = {bus.bus_id for bus in buses}
    line_ids = set()
    for row, record in _table(book["Lines"], ("line_id", "from_bus", "to_bus", "r_ohm", "x_ohm"),
                              ("ampacity_a", "status", "notes")):
        label = f"Lines row {row}"
        line = Line(
            _text(record["line_id"], f"{label} line_id"),
            _integer(record["from_bus"], f"{label} from_bus"),
            _integer(record["to_bus"], f"{label} to_bus"),
            _number(record["r_ohm"], f"{label} r_ohm"),
            _number(record["x_ohm"], f"{label} x_ohm"),
            _number(record["ampacity_a"], f"{label} ampacity_a") if record.get("ampacity_a") is not None else None,
        )
        if line.line_id in line_ids:
            raise ValueError(f"{label}: duplicate line_id {line.line_id}")
        line_ids.add(line.line_id)
        if line.from_bus not in bus_ids or line.to_bus not in bus_ids:
            raise ValueError(f"{label}: unknown endpoint")
        if line.from_bus == line.to_bus or line.r_ohm < 0 or (line.ampacity_a is not None and line.ampacity_a <= 0):
            raise ValueError(f"{label}: invalid line parameters")
        active = _flag(record["status"], f"{label} status") if "status" in record else True
        if active:
            lines.append(line)
    if "Loads" in book.sheetnames:
        for row, record in _table(book["Loads"], ("load_id", "bus_id", "p_base_kw", "q_base_kvar"), ("notes",)):
            label = f"Loads row {row}"
            loads.append(Load(_text(record["load_id"], f"{label} load_id"),
                              _integer(record["bus_id"], f"{label} bus_id"),
                              _number(record["p_base_kw"], f"{label} p_base_kw"),
                              _number(record["q_base_kvar"], f"{label} q_base_kvar")))
    if "PV" in book.sheetnames:
        for row, record in _table(book["PV"], ("pv_id", "bus_id", "p_rated_kw"),
                                  ("q_capability_kvar", "notes")):
            label = f"PV row {row}"
            pv_units.append(PVUnit(
                _text(record["pv_id"], f"{label} pv_id"),
                _integer(record["bus_id"], f"{label} bus_id"),
                _number(record["p_rated_kw"], f"{label} p_rated_kw"),
                _number(record["q_capability_kvar"], f"{label} q_capability_kvar")
                if record.get("q_capability_kvar") is not None else None,
            ))
    return NetworkModel(buses, lines, loads, pv_units, base, frequency_hz)


def _overview_quantity(sheet, label, unit):
    matches = [row[1] for row in sheet.iter_rows(values_only=True)
               if len(row) >= 2 and row[0] == label]
    if len(matches) != 1:
        raise ValueError(f"Overview: need exactly one {label} entry")
    value = _literal(matches[0], f"Overview {label}")
    if isinstance(value, str):
        match = re.match(r"^\s*(\d+(?:\.\d+)?)\s*" + re.escape(unit) + r"(?:\s|$)", value)
        if match is None:
            raise ValueError(f"Overview {label}: expected a value in {unit}")
        value = float(match.group(1))
    return _number(value, f"Overview {label}")


def _reference_network(book, frequency_hz):
    _check_sheets(book, ("Overview", "Bus Data", "Branch Data"),
                  ("Overview", "Bus Data", "Branch Data"))
    if frequency_hz is None:
        raise ValueError("Bus/branch reference workbook has no frequency: supply frequency_hz explicitly (e.g. 50 or 60 Hz)")
    base_kva = 1000.0 * _overview_quantity(book["Overview"], "Base MVA", "MVA")
    base_kv = _overview_quantity(book["Overview"], "Base voltage", "kV")
    buses, loads, lines = [], [], []
    demand = {}
    for row, record in _table(
            book["Bus Data"],
            ("Bus No.", "Type", "P_load (kW)", "Q_load (kVAr)", "Vm (p.u.) init", "Va (deg) init"),
            ("Notes",), reference_footer=True):
        label = f"Bus Data row {row}"
        bus_id = _integer(record["Bus No."], f"{label} Bus No.")
        kind = _text(record["Type"], f"{label} Type").lower()
        if kind not in ("slack (substation)", "pq (load)", "slack", "pq"):
            raise ValueError(f"{label}: unsupported bus type {kind}; only slack and PQ load buses are supported")
        voltage = _number(record["Vm (p.u.) init"], f"{label} voltage")
        if voltage <= 0 or _number(record["Va (deg) init"], f"{label} angle") != 0:
            raise ValueError(f"{label}: reference adapter needs positive initial voltage and zero angle")
        slack = kind.startswith("slack")
        buses.append(Bus(bus_id, f"Bus {bus_id}", base_kv, slack, voltage if slack else 1.0))
        p_kw = _number(record["P_load (kW)"], f"{label} P_load")
        q_kvar = _number(record["Q_load (kVAr)"], f"{label} Q_load")
        demand[bus_id] = (p_kw, q_kvar)
        if p_kw != 0 or q_kvar != 0:
            loads.append(Load(f"LD{bus_id}", bus_id, p_kw, q_kvar))
    for row, record in _table(
            book["Branch Data"],
            ("Branch No.", "From Bus", "To Bus", "R (Ohm)", "X (Ohm)",
             "P_load at To Bus (kW)", "Q_load at To Bus (kVAr)"),
            reference_footer=True):
        label = f"Branch Data row {row}"
        target = _integer(record["To Bus"], f"{label} To Bus")
        p_kw = _number(record["P_load at To Bus (kW)"], f"{label} receiving P_load")
        q_kvar = _number(record["Q_load at To Bus (kVAr)"], f"{label} receiving Q_load")
        if target not in demand or not all(
                isclose(a, b, rel_tol=1e-9, abs_tol=1e-9)
                for a, b in zip((p_kw, q_kvar), demand[target])):
            raise ValueError(f"{label}: receiving-bus load disagrees with Bus Data")
        lines.append(Line(
            f"L{_integer(record['Branch No.'], label + ' Branch No.')}",
            _integer(record["From Bus"], f"{label} From Bus"), target,
            _number(record["R (Ohm)"], f"{label} R"),
            _number(record["X (Ohm)"], f"{label} X"),
        ))
    return NetworkModel(buses, lines, loads, [], base_kva, frequency_hz)


def load_network(path: str | Path, *, workbook_format: str = "auto",
                 frequency_hz: float | None = None) -> NetworkModel:
    """Import a strict project template or supported bus/branch workbook."""
    source = Path(path)
    if source.suffix.lower() != ".xlsx":
        raise ValueError("Network import supports .xlsx only; renaming another format does not convert it")
    if workbook_format not in ("auto", "project", "reference", "ieee33"):
        raise ValueError("workbook_format must be auto, project, reference or ieee33")
    if frequency_hz is not None:
        frequency_hz = _number(frequency_hz, "frequency_hz")
        if frequency_hz <= 0:
            raise ValueError("frequency_hz must be positive")
    try:
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, InvalidFileException) as error:
        raise ValueError(f"Not a readable .xlsx workbook: {source}") from error
    try:
        if "Generator Data" in book.sheetnames:
            raise ValueError("PJM/generator workbook is unsupported: our model needs a radial feeder with PQ loads, not meshed transmission/PV generator buses")
        selected = workbook_format
        if selected == "auto":
            if {"Network", "Buses", "Lines"}.issubset(book.sheetnames):
                selected = "project"
            elif {"Overview", "Bus Data", "Branch Data"}.issubset(book.sheetnames):
                selected = "reference"
            else:
                raise ValueError("Unrecognized workbook layout; use the documented project template or bus/branch reference workbook")
        network = _project_network(book, frequency_hz) if selected == "project" else _reference_network(book, frequency_hz)
        if selected == "ieee33" and (len(network.buses) != 33 or len(network.lines) != 32):
            raise ValueError(
                f"IEEE33 format expects 33 buses and 32 branches, but this workbook has "
                f"{len(network.buses)} buses and {len(network.lines)} branches; "
                "choose Auto or Bus / Branch instead"
            )
        network.validate()
        return network
    finally:
        book.close()
