"""Static workbook schemas, incompatible cases and real IEEE33 integration."""
import copy
from hashlib import sha256
from pathlib import Path
import sys
import unittest
from unittest.mock import patch
from zipfile import BadZipFile

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from src.fbs import FBSConfig, run_fbs
from src.main import base_operating_point
from src.network_loader import load_network


DATA = Path(__file__).resolve().parents[2] / "04_Research/Working_Data/Network_Data"
IEEE33 = DATA / "IEEE_33Bus_Network_Data.xlsx"
PJM = DATA / "PJM_5Bus_Network_Data.xlsx"


class Sheet:
    def __init__(self, title, rows):
        self.title = title
        self.rows = rows

    def iter_rows(self, values_only=False):
        return iter(self.rows)


class Workbook:
    """In-memory input fixture: no source workbook is edited or authored."""
    def __init__(self, tables):
        self.sheetnames = list(tables)
        self.sheets = {name: Sheet(name, rows) for name, rows in tables.items()}
        self.closed = False

    def __getitem__(self, name):
        return self.sheets[name]

    def close(self):
        self.closed = True


def project_tables():
    return {
        "Network": [["parameter", "value"], ["base_power_kva", 1000], ["frequency_hz", 60]],
        "Buses": [["bus_id", "base_kv", "is_slack"], [10, 11, False], [20, 11, True], [35, 11, False]],
        "Lines": [["line_id", "from_bus", "to_bus", "r_ohm", "x_ohm"],
                  ["L1", 10, 20, 0.1, 0.2], ["L2", 35, 10, 0.15, 0.3]],
        "Loads": [["load_id", "bus_id", "p_base_kw", "q_base_kvar"], ["LD35", 35, 100, 50]],
        "PV": [["pv_id", "bus_id", "p_rated_kw"], ["PV10", 10, 25]],
    }


def reference_tables():
    book = load_workbook(IEEE33, read_only=True, data_only=False)
    try:
        return {sheet.title: [list(row) for row in sheet.values] for sheet in book.worksheets}
    finally:
        book.close()


class NetworkLoaderTests(unittest.TestCase):
    def load_tables(self, tables, **kwargs):
        book = Workbook(tables)
        with patch("src.network_loader.load_workbook", return_value=book) as reader:
            try:
                network = load_network("fixture.xlsx", **kwargs)
                reader.assert_called_once_with(Path("fixture.xlsx"), read_only=True, data_only=False, keep_links=False)
                return network
            finally:
                self.assertTrue(book.closed)

    def test_project_template_noncontiguous_ids_and_reversed_lines(self):
        network = self.load_tables(project_tables())
        self.assertEqual(network.slack_bus_id(), 20)
        self.assertEqual(network.base_power_kva, 1000)
        self.assertEqual(network.frequency_hz, 60)
        self.assertEqual(network.radial_order().parent_bus_by_bus, {10: 20, 35: 10})
        self.assertEqual(network.loads[0].p_base_kw, 100)
        self.assertEqual(network.pv_units[0].p_rated_kw, 25)
        self.assertEqual(network.lines[0].r_ohm, 0.1)

    def test_optional_fields_names_voltage_ampacity_and_q_capability(self):
        tables = project_tables()
        tables["Buses"][0] += ["name", "v_set_pu"]
        for row in tables["Buses"][1:]:
            row += [f"Name {row[0]}", 1.02]
        tables["Lines"][0].append("ampacity_a")
        for row in tables["Lines"][1:]:
            row.append(250)
        tables["PV"][0].append("q_capability_kvar")
        tables["PV"][1].append(10)
        network = self.load_tables(tables)
        self.assertEqual(network.buses[0].name, "Name 10")
        self.assertEqual(network.buses[1].v_set_pu, 1.02)
        self.assertEqual(network.lines[0].ampacity_a, 250)
        self.assertEqual(network.pv_units[0].q_capability_kvar, 10)

    def test_optional_equipment_sheets_can_be_absent(self):
        tables = project_tables()
        del tables["Loads"], tables["PV"]
        network = self.load_tables(tables)
        self.assertEqual(network.loads, [])
        self.assertEqual(network.pv_units, [])

    def test_explicit_frequency_override_and_missing_frequency(self):
        tables = project_tables()
        self.assertEqual(self.load_tables(tables, frequency_hz=50).frequency_hz, 50)
        tables["Network"].pop()
        with self.assertRaisesRegex(ValueError, "frequency_hz"):
            self.load_tables(tables)
        self.assertEqual(self.load_tables(tables, frequency_hz=60).frequency_hz, 60)

    def test_unsupported_or_duplicate_network_parameters(self):
        for row in (["frequency_hz", 50], ["transformer_tap", 1.1]):
            tables = project_tables()
            tables["Network"].append(row)
            with self.subTest(row=row), self.assertRaises(ValueError):
                self.load_tables(tables)

    def test_open_tie_is_omitted_but_closed_tie_is_not_silently_removed(self):
        tables = project_tables()
        tables["Lines"][0].append("status")
        for row in tables["Lines"][1:]:
            row.append(1)
        tables["Lines"].append(["tie", 20, 35, 0.2, 0.1, 0])
        self.assertEqual(len(self.load_tables(tables).lines), 2)
        tables["Lines"][-1][-1] = 1
        with self.assertRaisesRegex(ValueError, "radial"):
            self.load_tables(tables)

    def test_open_lines_still_need_valid_endpoints_parameters_and_unique_ids(self):
        for tie in (["L1", 20, 35, 0.2, 0.1, 0], ["tie", 20, 999, 0.2, 0.1, 0],
                    ["tie", 20, 35, -1, 0.1, 0]):
            tables = project_tables()
            tables["Lines"][0].append("status")
            for row in tables["Lines"][1:]:
                row.append(1)
            tables["Lines"].append(tie)
            with self.subTest(tie=tie), self.assertRaises(ValueError):
                self.load_tables(tables)

    def test_unknown_sheet_or_columns_rejected(self):
        for defect in ("sheet", "column", "missing", "duplicate"):
            tables = project_tables()
            if defect == "sheet":
                tables["Transformers"] = []
            elif defect == "column":
                tables["Lines"][0].append("shunt_b_pu")
            elif defect == "missing":
                tables["Lines"][0][3] = "wrong_name"
            else:
                tables["Lines"][0].append("r_ohm")
            with self.subTest(defect=defect), self.assertRaises(ValueError):
                self.load_tables(tables)

    def test_data_formulas_strings_missing_and_nonfinite_numbers_rejected(self):
        for value in ("=0.1+0.1", "0.1", None, float("nan"), float("inf"), True):
            tables = project_tables()
            tables["Lines"][1][3] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.load_tables(tables)

    def test_bad_ids_flags_and_equipment_references(self):
        for defect in ("fractional_id", "duplicate_bus", "two_slacks", "no_slack", "load_bus", "pv_bus", "flag"):
            tables = project_tables()
            if defect == "fractional_id":
                tables["Buses"][1][0] = 10.5
            elif defect == "duplicate_bus":
                tables["Buses"].append(tables["Buses"][1][:])
            elif defect == "two_slacks":
                tables["Buses"][1][2] = True
            elif defect == "no_slack":
                tables["Buses"][2][2] = False
            elif defect == "load_bus":
                tables["Loads"][1][1] = 999
            elif defect == "pv_bus":
                tables["PV"][1][1] = 999
            else:
                tables["Buses"][1][2] = "yes"
            with self.subTest(defect=defect), self.assertRaises(ValueError):
                self.load_tables(tables)

    def test_disconnected_multivoltage_and_negative_load_rejected(self):
        for defect in ("disconnected", "multivoltage", "negative_load"):
            tables = project_tables()
            if defect == "disconnected":
                tables["Lines"].pop()
            elif defect == "multivoltage":
                tables["Buses"][1][1] = 0.4
            else:
                tables["Loads"][1][2] = -10
            with self.subTest(defect=defect), self.assertRaises(ValueError):
                self.load_tables(tables)

    def test_actual_ieee33_workbook_and_base_load_snapshot(self):
        before = sha256(IEEE33.read_bytes()).hexdigest()
        network = load_network(IEEE33, frequency_hz=60)
        self.assertEqual(len(network.buses), 33)
        self.assertEqual(len(network.lines), 32)
        self.assertEqual(network.slack_bus_id(), 1)
        self.assertEqual(network.base_power_kva, 100000)
        self.assertEqual(network.buses[0].base_kv, 12.66)
        self.assertEqual(sum(load.p_base_kw for load in network.loads), 3715)
        self.assertEqual(sum(load.q_base_kvar for load in network.loads), 2300)
        self.assertEqual(network.pv_units, [])
        self.assertEqual(len(network.loads), 32)  # No invented zero-load device at the slack bus.
        self.assertEqual(network.lines[0].r_ohm, 0.0922)
        flow = run_fbs(network, base_operating_point(network), FBSConfig())
        self.assertTrue(flow.converged)
        self.assertAlmostEqual(flow.total_loss_p_kw, 202.67822, delta=0.001)
        self.assertAlmostEqual(min(flow.voltage_mag_pu.values()), 0.91308, delta=1e-5)
        self.assertEqual(sha256(IEEE33.read_bytes()).hexdigest(), before)

    def test_imported_network_uses_absolute_hourly_load_without_adding_base_load(self):
        from io import StringIO
        from src.data_loader import ProfileData
        from src.simulation import SimulationConfig, simulate_timeseries
        network = load_network(IEEE33, frequency_hz=60)
        load_rows = ["timestamp,bus_id,p_kw,q_kvar"]
        pv_rows = ["timestamp,bus_id,p_kw,q_kvar"]
        for hour in (0, 1):
            timestamp = f"2026-01-01T0{hour}:00:00"
            for load in network.loads:
                load_rows.append(f"{timestamp},{load.bus_id},{load.p_base_kw * 0.5},{load.q_base_kvar * 0.5}")
            pv_rows.append(f"{timestamp},1,0,0")
        profiles = ProfileData(StringIO("\n".join(load_rows)), StringIO("\n".join(pv_rows)))
        result = simulate_timeseries(network, profiles, SimulationConfig())
        self.assertTrue(result.completed and result.converged)
        self.assertEqual(sum(result.steps[0].operating_point.p_load_kw.values()), 3715 * 0.5)
        self.assertEqual(sum(result.steps[0].operating_point.q_load_kvar.values()), 2300 * 0.5)

    def test_reference_requires_explicit_frequency(self):
        with self.assertRaisesRegex(ValueError, "supply frequency_hz"):
            load_network(IEEE33)

    def test_pjm_rejected_without_changing_it(self):
        before = sha256(PJM.read_bytes()).hexdigest()
        with self.assertRaisesRegex(ValueError, "PJM/generator"):
            load_network(PJM, frequency_hz=60)
        self.assertEqual(sha256(PJM.read_bytes()).hexdigest(), before)

    def test_reference_overview_units_and_literal_values(self):
        for value in ("100 kVA", "=100", None, "unknown"):
            tables = reference_tables()
            tables["Overview"][7][1] = value
            with self.subTest(value=value), self.assertRaises(ValueError):
                self.load_tables(tables, frequency_hz=60)

    def test_reference_receiving_loads_must_agree_not_double_count(self):
        tables = reference_tables()
        tables["Branch Data"][1][5] = 999
        with self.assertRaisesRegex(ValueError, "disagrees"):
            self.load_tables(tables, frequency_hz=60)

    def test_reference_generator_types_bad_angles_and_missing_rows(self):
        for defect in ("generator", "angle", "missing_row"):
            tables = reference_tables()
            if defect == "generator":
                tables["Bus Data"][2][1] = "PV (generator)"
            elif defect == "angle":
                tables["Bus Data"][1][5] = 20
            else:
                tables["Branch Data"].pop(32)
            with self.subTest(defect=defect), self.assertRaises(ValueError):
                self.load_tables(tables, frequency_hz=60)

    def test_reference_total_formulas_ignored_but_data_formulas_rejected(self):
        tables = reference_tables()
        self.assertTrue(tables["Bus Data"][-1][2].startswith("="))
        self.assertEqual(len(self.load_tables(tables, frequency_hz=60).buses), 33)
        tables["Bus Data"][2][2] = "=100"
        with self.assertRaisesRegex(ValueError, "literal value"):
            self.load_tables(tables, frequency_hz=60)

    def test_reference_records_cannot_be_hidden_after_footer(self):
        tables = reference_tables()
        tables["Bus Data"].append(tables["Bus Data"][1][:])
        with self.assertRaisesRegex(ValueError, "after the footer"):
            self.load_tables(tables, frequency_hz=60)

    def test_filename_does_not_select_schema_and_inputs_not_mutated(self):
        tables = project_tables()
        before = copy.deepcopy(tables)
        self.load_tables(tables, workbook_format="project")
        self.assertEqual(tables, before)
        with self.assertRaises(ValueError):
            self.load_tables(tables, workbook_format="ieee33", frequency_hz=60)

    def test_invalid_extension_format_frequency_missing_file_and_corrupt_file(self):
        for filename in ("case.csv", "case.mat", "case.m", "case.xls", "case.xlsm"):
            with self.subTest(filename=filename), self.assertRaisesRegex(ValueError, ".xlsx only"):
                load_network(filename)
        for value in (0, -1, True, "60", float("nan")):
            with self.subTest(frequency=value), self.assertRaises(ValueError):
                load_network(IEEE33, frequency_hz=value)
        with self.assertRaises(ValueError):
            load_network(IEEE33, workbook_format="guess")
        with self.assertRaises(FileNotFoundError):
            load_network(DATA / "not_present.xlsx")
        with patch("src.network_loader.load_workbook", side_effect=BadZipFile("invalid")):
            with self.assertRaisesRegex(ValueError, "Not a readable"):
                load_network("invalid.xlsx")


if __name__ == "__main__":
    unittest.main()
